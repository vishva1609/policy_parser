"""
compliance_engine/exporter.py
──────────────────────────────
Step 8 (partial): Export comparison results to PDF and Excel.

PDF  → uses reportlab (pure Python, no system dependencies)
Excel → uses openpyxl
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

# ─────────────────────────────────────────────────────────────────────────────
#  Excel export
# ─────────────────────────────────────────────────────────────────────────────

def export_excel(report: dict, out_path: str) -> str:
    """Write the compliance report to an Excel workbook with multiple sheets."""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment
    except ImportError:
        raise ImportError("Install openpyxl: pip install openpyxl")

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"
    _fill_summary_sheet(ws_sum, report)

    # ── Sheet 2: All Clauses ─────────────────────────────────────────────────
    ws_clauses = wb.create_sheet("Clause Detail")
    _fill_clauses_sheet(ws_clauses, report.get("clause_details", []))

    # ── Sheet 3: Critical Gaps ───────────────────────────────────────────────
    ws_gaps = wb.create_sheet("Critical Gaps")
    _fill_gap_sheet(ws_gaps, report.get("critical_gaps", []), "Critical Gaps")

    # ── Sheet 4: Missing Clauses ─────────────────────────────────────────────
    ws_missing = wb.create_sheet("Missing Clauses")
    _fill_gap_sheet(ws_missing, report.get("missing_clauses", []), "Missing Clauses")

    # ── Sheet 5: Category Scores ─────────────────────────────────────────────
    ws_cat = wb.create_sheet("Category Scores")
    _fill_category_sheet(ws_cat, report.get("category_scores", {}))

    wb.save(out_path)
    return out_path


def _header_row(ws, headers: list[str], fill_color: str = "1E3A5F"):
    from openpyxl.styles import PatternFill, Font, Alignment
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _status_fill(status: str):
    from openpyxl.styles import PatternFill
    colors = {"Yes": "C6EFCE", "Partial": "FFEB9C", "No": "FFC7CE"}
    c = colors.get(status, "FFFFFF")
    return PatternFill(start_color=c, end_color=c, fill_type="solid")


def _fill_summary_sheet(ws, report: dict):
    from openpyxl.styles import Font
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    rows = [
        ("Overall Compliance %",      report.get("overall_compliance_pct", 0)),
        ("Total Clauses Compared",    report.get("total_clauses", 0)),
        ("Compliant (Yes)",           report.get("compliant", 0)),
        ("Partial Matches",           report.get("partial_matches", 0)),
        ("Critical Gaps",             report.get("gaps", 0)),
        ("Missing Clauses",           report.get("missing", 0)),
    ]
    for i, (label, value) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)


def _fill_clauses_sheet(ws, clauses: list[dict]):
    headers = ["Clause A ID", "Clause A Title", "Category", "Strength",
               "Status", "Best Match ID", "Best Match Title", "Similarity",
               "Gap Type", "Reason", "Gap"]
    _header_row(ws, headers)
    for r, c in enumerate(clauses, 2):
        vals = [
            c.get("clause_a_id"), c.get("clause_a_title"), c.get("category"),
            c.get("strength"), c.get("status"), c.get("best_match_id"),
            c.get("best_match_title"), c.get("best_similarity"),
            c.get("gap_type"), c.get("reason"), c.get("gap"),
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=v)
            if col == 5:  # Status
                cell.fill = _status_fill(str(v))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22


def _fill_gap_sheet(ws, items: list[dict], title: str):
    headers = ["Clause A ID", "Title", "Status", "Strength", "Category",
               "Gap", "Reason", "Confidence"]
    _header_row(ws, headers)
    for r, item in enumerate(items, 2):
        vals = [
            item.get("clause_a_id"), item.get("clause_a_title"),
            item.get("status"), item.get("strength"), item.get("category"),
            item.get("gap"), item.get("reason"), item.get("confidence"),
        ]
        for col, v in enumerate(vals, 1):
            ws.cell(row=r, column=col, value=v)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22


def _fill_category_sheet(ws, cat_scores: dict):
    _header_row(ws, ["Category", "Compliance %"])
    for r, (cat, data) in enumerate(cat_scores.items(), 2):
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=data.get("compliance_pct", 0))


# ─────────────────────────────────────────────────────────────────────────────
#  PDF export
# ─────────────────────────────────────────────────────────────────────────────

def export_pdf(report: dict, policy_a_name: str, policy_b_name: str,
               out_path: str) -> str:
    """Write a multi-page compliance PDF report."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        )
    except ImportError:
        raise ImportError("Install reportlab: pip install reportlab")

    doc = SimpleDocTemplate(out_path, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story = []

    # ── Title ────────────────────────────────────────────────────────────────
    title_style = ParagraphStyle("title", parent=styles["Title"],
                                 fontSize=20, spaceAfter=12)
    story.append(Paragraph("Policy Compliance Report", title_style))
    story.append(Paragraph(
        f"<b>Policy A:</b> {policy_a_name} &nbsp;&nbsp; vs &nbsp;&nbsp; "
        f"<b>Policy B:</b> {policy_b_name}", styles["Normal"]))
    story.append(Spacer(1, 0.5*cm))

    # ── Summary table ────────────────────────────────────────────────────────
    pct = report.get("overall_compliance_pct", 0)
    summary_data = [
        ["Metric", "Value"],
        ["Overall Compliance", f"{pct:.1f}%"],
        ["Total Clauses Compared", str(report.get("total_clauses", 0))],
        ["Compliant (Yes)", str(report.get("compliant", 0))],
        ["Partial Matches", str(report.get("partial_matches", 0))],
        ["Critical Gaps", str(report.get("gaps", 0))],
        ["Missing Clauses", str(report.get("missing", 0))],
    ]
    t = Table(summary_data, colWidths=[8*cm, 6*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
        ("GRID",       (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7FF")]),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    # ── Category scores ──────────────────────────────────────────────────────
    cat_scores = report.get("category_scores", {})
    if cat_scores:
        story.append(Paragraph("Category Compliance Scores", styles["Heading2"]))
        cat_data = [["Category", "Compliance %"]] + [
            [k, f"{v.get('compliance_pct', 0):.1f}%"]
            for k, v in cat_scores.items()
        ]
        ct = Table(cat_data, colWidths=[8*cm, 6*cm])
        ct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E6DA4")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID",       (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EBF5FF")]),
        ]))
        story.append(ct)
        story.append(Spacer(1, 0.4*cm))
        story.append(PageBreak())

    # ── Critical Gaps ────────────────────────────────────────────────────────
    gaps = report.get("critical_gaps", [])
    if gaps:
        story.append(Paragraph("Critical Gaps", styles["Heading2"]))
        for g in gaps:
            story.append(Paragraph(
                f"<b>{g.get('clause_a_id')} – {g.get('clause_a_title')}</b>  "
                f"[{g.get('strength', '').upper()}]",
                styles["Normal"]))
            story.append(Paragraph(
                f"<i>Gap:</i> {g.get('gap', '')}", styles["Normal"]))
            story.append(Spacer(1, 0.2*cm))
        story.append(PageBreak())

    # ── All Clauses table ─────────────────────────────────────────────────────
    story.append(Paragraph("Clause-by-Clause Comparison", styles["Heading2"]))
    col_headers = ["Clause A", "Status", "Category", "Best Match B", "Similarity"]
    table_data  = [col_headers]
    _STATUS_COLORS = {
        "Yes": colors.HexColor("#C6EFCE"),
        "Partial": colors.HexColor("#FFEB9C"),
        "No": colors.HexColor("#FFC7CE"),
    }
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 7),
        ("GRID",       (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN",     (0, 0), (-1, -1), "TOP"),
        ("WORDWRAP",   (0, 0), (-1, -1), True),
    ]
    for row_i, c in enumerate(report.get("clause_details", []), 2):
        status = c.get("status", "")
        row = [
            f"{c.get('clause_a_id', '')} {c.get('clause_a_title', '')}",
            status,
            c.get("category", ""),
            f"{c.get('best_match_id', '')} {c.get('best_match_title', '')}",
            f"{c.get('best_similarity', 0):.2f}",
        ]
        table_data.append(row)
        sc = _STATUS_COLORS.get(status, colors.white)
        style_cmds.append(("BACKGROUND", (1, row_i - 1), (1, row_i - 1), sc))

    clause_table = Table(
        table_data,
        colWidths=[5*cm, 1.5*cm, 2.5*cm, 5*cm, 2*cm],
        repeatRows=1,
    )
    clause_table.setStyle(TableStyle(style_cmds))
    story.append(clause_table)

    doc.build(story)
    return out_path
